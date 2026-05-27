from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from calculate_metrics import DEFAULT_OUTPUT as DEFAULT_METRICS_PATH, calculate_metrics
from clean_data import DEFAULT_OUTPUT as DEFAULT_CLEAN_PATH, load_csv
from detect_anomalies import DEFAULT_OUTPUT as DEFAULT_ANOMALIES_PATH
from generate_sample_data import DEFAULT_OUTPUT as DEFAULT_RAW_PATH, PROJECT_ROOT


DEFAULT_OUTPUT = PROJECT_ROOT / "reports" / "adops_dashboard.xlsx"

SUMMARY_GROUPS = {
    "Summary by date": ["date"],
    "Summary by site_section": ["site_section"],
    "Summary by device": ["device"],
    "Summary by ad_unit": ["ad_unit"],
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export the AdOps dashboard workbook.")
    parser.add_argument("--raw", type=Path, default=DEFAULT_RAW_PATH, help="Raw CSV input path.")
    parser.add_argument("--clean", type=Path, default=DEFAULT_CLEAN_PATH, help="Clean CSV input path.")
    parser.add_argument("--metrics", type=Path, default=DEFAULT_METRICS_PATH, help="Metrics CSV input path.")
    parser.add_argument(
        "--anomalies",
        type=Path,
        default=DEFAULT_ANOMALIES_PATH,
        help="Anomalies CSV input path.",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Excel output path.")
    return parser.parse_args()


def aggregate_summary(metrics_df: pd.DataFrame, group_columns: list[str]) -> pd.DataFrame:
    numeric_totals = {
        "ad_requests": "sum",
        "matched_requests": "sum",
        "impressions": "sum",
        "clicks": "sum",
        "estimated_revenue": "sum",
    }
    grouped = metrics_df.groupby(group_columns, dropna=False).agg(numeric_totals).reset_index()
    grouped = calculate_metrics(grouped)

    ordered_columns = group_columns + [
        "ad_requests",
        "matched_requests",
        "impressions",
        "clicks",
        "estimated_revenue",
        "fill_rate",
        "ctr",
        "rpm",
        "ecpm",
        "match_rate",
    ]
    return grouped[ordered_columns]


def build_summary_tables(metrics_df: pd.DataFrame) -> list[tuple[str, pd.DataFrame]]:
    return [
        (title, aggregate_summary(metrics_df, group_columns))
        for title, group_columns in SUMMARY_GROUPS.items()
    ]


def write_summary_sheet(
    writer: pd.ExcelWriter,
    summary_tables: list[tuple[str, pd.DataFrame]],
    sheet_name: str = "summary",
) -> None:
    start_row = 0
    for title, table in summary_tables:
        table.to_excel(writer, sheet_name=sheet_name, startrow=start_row + 1, index=False)
        worksheet = writer.sheets[sheet_name]
        worksheet.cell(row=start_row + 1, column=1, value=title)
        start_row += len(table) + 4


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0

        for cell in column_cells:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))

        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)


def format_workbook(writer: pd.ExcelWriter) -> None:
    workbook = writer.book
    header_fill = PatternFill("solid", fgColor="1F4E78")
    title_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        autosize_columns(worksheet)

        for row in worksheet.iter_rows():
            first_value = row[0].value
            if first_value in SUMMARY_GROUPS:
                row[0].font = title_font
                row[0].fill = title_fill
                continue

            if any(cell.value in ["date", "site_section", "device", "country", "ad_unit"] for cell in row):
                for cell in row:
                    if cell.value is not None:
                        cell.fill = header_fill
                        cell.font = white_font


def export_dashboard(
    raw_df: pd.DataFrame,
    clean_df: pd.DataFrame,
    metrics_df: pd.DataFrame,
    anomalies_df: pd.DataFrame,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    summary_tables = build_summary_tables(metrics_df)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        raw_df.to_excel(writer, sheet_name="raw_data", index=False)
        clean_df.to_excel(writer, sheet_name="clean_data", index=False)
        metrics_df.to_excel(writer, sheet_name="metrics", index=False)
        anomalies_df.to_excel(writer, sheet_name="anomalies", index=False)
        write_summary_sheet(writer, summary_tables)
        format_workbook(writer)


def main() -> None:
    args = parse_args()

    try:
        raw_df = load_csv(args.raw)
        clean_df = load_csv(args.clean)
        metrics_df = load_csv(args.metrics)
        anomalies_df = load_csv(args.anomalies)
        export_dashboard(raw_df, clean_df, metrics_df, anomalies_df, args.output)
    except Exception as exc:
        raise SystemExit(f"Failed to export dashboard: {exc}") from exc

    print(f"Exported dashboard at {args.output}")


if __name__ == "__main__":
    main()
